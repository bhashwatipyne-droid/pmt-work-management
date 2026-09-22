"""Read and validate a CSV / Excel sheet of deliverables for the bulk-import feature.

Pure functions: no database, no web framework. server.py loads the sheet with
read_table(), checks every row with validate_table(), and only then writes the
rows that came out as "ok".

Columns (header names are matched loosely - case, spaces and common synonyms are
ignored). Only Name is required.

    Name | Type | Start date | End date | Stages | Approvals
"""

import csv
import difflib
import io
import re
from datetime import date, datetime, timedelta
from typing import Any, Dict, Iterable, List, Optional, Set, Tuple

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_ROWS = 500
MAX_NAME_LENGTH = 200


class ImportFileError(ValueError):
    """The file as a whole cannot be used (wrong format, no header, too big...)."""


HEADER_ALIASES = {
    "name": {"name", "deliverable", "deliverable name", "deliverable title", "title"},
    "type": {"type", "deliverable type", "format", "activity", "activity type"},
    "start": {"start", "start date", "start dt", "startdt", "from", "starts"},
    "end": {"end", "end date", "end dt", "enddt", "due", "due date", "deadline", "to", "ends"},
    "stages": {"stages", "stage", "required stages", "production stages", "teams", "team"},
    "approvals": {"approvals", "approval", "approval types", "approval type", "additional approvals", "approvers"},
}

STAGE_ALIASES = {"animation": "Animate", "animate": "Animate", "content": "Content", "design": "Design"}

APPROVAL_ALIASES = {
    "manager": None,            # always present, never needs listing
    "leadership": "LEADERSHIP",
    "clientspoc": "CLIENT_SPOC",
    "spoc": "CLIENT_SPOC",
    "client": "CLIENT_SPOC",
    "compliance": "COMPLIANCE",
}

DATE_FORMATS = (
    "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%d-%m-%y",
    "%d %b %Y", "%d %B %Y", "%b %d, %Y", "%B %d, %Y", "%Y/%m/%d",
)


# ---------------------------------------------------------------- reading

def read_table(filename: str, content: bytes) -> List[List[Any]]:
    """The sheet as a list of rows (each a list of cell values)."""
    if len(content) > MAX_FILE_BYTES:
        raise ImportFileError("The file is larger than 2 MB. Split it into smaller files.")
    if not content:
        raise ImportFileError("The file is empty.")

    extension = (filename or "").lower().rsplit(".", 1)[-1] if "." in (filename or "") else ""

    if extension == "csv":
        return _read_csv(content)
    if extension == "xlsx":
        return _read_xlsx(content)
    if extension == "xls":
        raise ImportFileError("Old .xls files are not supported. Save the sheet as .xlsx or .csv and upload that.")
    raise ImportFileError("Only .csv and .xlsx files can be imported.")


def _read_csv(content: bytes) -> List[List[Any]]:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ImportFileError("The file could not be read as text. Save it as CSV (UTF-8) and try again.")

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [list(row) for row in csv.reader(io.StringIO(text), dialect)]


def _read_xlsx(content: bytes) -> List[List[Any]]:
    try:
        from openpyxl import load_workbook  # imported here so a missing package can never stop the server booting
    except ImportError:
        raise ImportFileError("Excel import is not available on the server yet. Upload the sheet as a .csv file instead.")

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise ImportFileError("This does not look like a valid .xlsx file.")

    sheet = workbook.worksheets[0] if workbook.worksheets else None
    if sheet is None:
        raise ImportFileError("The workbook has no sheets.")
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    workbook.close()
    return rows


# ---------------------------------------------------------------- cell helpers

def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _key(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def _loose(text: str) -> str:
    """Comparison key that ignores case, punctuation and a trailing plural 's'."""
    key = _key(text)
    return key[:-1] if key.endswith("s") and len(key) > 3 else key


def parse_date(value: Any) -> Optional[str]:
    """YYYY-MM-DD from a date cell, an Excel serial or text (day first for 12/03/2026).
    None for a blank cell; ValueError for something that is not a date."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if 20000 <= float(value) <= 80000:            # an Excel serial number
            return (date(1899, 12, 30) + timedelta(days=int(value))).isoformat()
        raise ValueError("not a date")
    text = str(value).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError("not a date")


def _suggest_types(raw: str, canonical: Iterable[str], limit: int = 2) -> List[str]:
    """The closest known types to something mistyped. Scores the whole name and each
    word of it, so "Reelz" finds "Reel / Short Video"."""
    wanted = _key(raw)
    scored = []
    for name in canonical:
        key = _key(name)
        whole = difflib.SequenceMatcher(None, wanted, key).ratio()
        best_word = max((difflib.SequenceMatcher(None, wanted, word).ratio() for word in key.split()), default=0.0)
        score = max(whole, best_word * 0.95)
        if score >= 0.75:
            scored.append((score, name))
    scored.sort(key=lambda pair: (-pair[0], pair[1]))
    return [name for _, name in scored[:limit]]


def _split_list(text: str) -> List[str]:
    return [part.strip() for part in re.split(r"[,;|/\n]+", text) if part.strip()]


# ---------------------------------------------------------------- validation

def _detect_columns(header: List[Any]) -> Tuple[Dict[str, int], List[str]]:
    columns: Dict[str, int] = {}
    ignored: List[str] = []
    for index, cell in enumerate(header):
        label = _text(cell)
        if not label:
            continue
        key = _key(label)
        for field, aliases in HEADER_ALIASES.items():
            if key in aliases and field not in columns:
                columns[field] = index
                break
        else:
            ignored.append(label)
    return columns, ignored


def validate_table(
    table: List[List[Any]],
    deliverable_types: Iterable[str],
    stages: Iterable[str],
    approval_types: Iterable[str],
    existing: Set[Tuple[str, str]],
) -> Dict[str, Any]:
    """Check every data row. Nothing is written here.

    `existing` holds (name, type) pairs - lower-cased - already in the project, so
    uploading the same sheet twice does not create every deliverable twice."""
    header_index = next(
        (i for i, row in enumerate(table) if any(_text(c) for c in row)),
        None,
    )
    if header_index is None:
        raise ImportFileError("The file has no rows.")

    columns, ignored = _detect_columns(table[header_index])
    if "name" not in columns:
        raise ImportFileError(
            "Could not find a Name column. The first row must be a header like: "
            "Name, Type, Start date, End date, Stages, Approvals."
        )

    body = [
        (i + 1, row) for i, row in enumerate(table[header_index + 1:], start=header_index + 1)
        if any(_text(c) for c in row)
    ]
    if not body:
        raise ImportFileError("The file has a header but no deliverables under it.")
    if len(body) > MAX_ROWS:
        raise ImportFileError(f"The file has {len(body)} rows. The limit is {MAX_ROWS} per upload.")

    type_by_key = {_key(t): t for t in deliverable_types}
    type_by_loose = {}
    for t in type_by_key.values():
        type_by_loose.setdefault(_loose(t), t)
    valid_stages = list(stages)
    valid_approvals = set(approval_types)

    def cell(row, field):
        index = columns.get(field)
        return row[index] if index is not None and index < len(row) else None

    seen: Dict[Tuple[str, str], int] = {}
    results: List[Dict[str, Any]] = []

    for sheet_row, row in body:
        errors: List[str] = []
        warnings: List[str] = []

        name = _text(cell(row, "name"))
        if not name:
            errors.append("Name is empty")
        elif len(name) > MAX_NAME_LENGTH:
            errors.append(f"Name is longer than {MAX_NAME_LENGTH} characters")

        # ---- type
        type_raw = _text(cell(row, "type"))
        type_value = ""
        if type_raw:
            match = type_by_key.get(_key(type_raw))
            if match is None:
                loose = type_by_loose.get(_loose(type_raw))
                if loose:
                    match = loose
                    warnings.append(f'"{type_raw}" matched to "{loose}"')
            if match is None:
                hints = _suggest_types(type_raw, type_by_key.values())
                errors.append(
                    f'Unknown type "{type_raw}"'
                    + (". Did you mean " + " or ".join(f'"{h}"' for h in hints) + "?" if hints else ". Use a type from the list in the template.")
                )
            else:
                type_value = match
        else:
            warnings.append("No type: time cannot be auto-filled and the work category stays blank")

        # ---- dates
        start_dt = end_dt = None
        try:
            start_dt = parse_date(cell(row, "start"))
        except ValueError:
            errors.append(f'Start date "{_text(cell(row, "start"))}" is not a valid date (use YYYY-MM-DD or DD/MM/YYYY)')
        try:
            end_dt = parse_date(cell(row, "end"))
        except ValueError:
            errors.append(f'End date "{_text(cell(row, "end"))}" is not a valid date (use YYYY-MM-DD or DD/MM/YYYY)')
        if start_dt and end_dt and end_dt < start_dt:
            errors.append("End date is before the start date")

        # ---- stages
        stage_raw = _text(cell(row, "stages"))
        chosen_stages: List[str] = []
        if stage_raw:
            for token in _split_list(stage_raw):
                mapped = STAGE_ALIASES.get(_key(token))
                if mapped is None or mapped not in valid_stages:
                    errors.append(f'Unknown stage "{token}" (use Content, Design, Animate)')
                elif mapped not in chosen_stages:
                    chosen_stages.append(mapped)
        chosen_stages = [s for s in valid_stages if s in chosen_stages] or (["Content"] if not stage_raw else [])
        if not stage_raw:
            warnings.append("No stages: defaulted to Content")

        # ---- approvals
        approvals: List[str] = []
        for token in _split_list(_text(cell(row, "approvals"))):
            key = re.sub(r"[^a-z]", "", token.lower())
            if key not in APPROVAL_ALIASES:
                errors.append(f'Unknown approval "{token}" (use Leadership, Client SPOC, Compliance)')
                continue
            mapped = APPROVAL_ALIASES[key]
            if mapped and mapped in valid_approvals and mapped not in approvals:
                approvals.append(mapped)

        # ---- duplicates
        status = "error" if errors else "ok"
        if status == "ok":
            dup_key = (name.lower(), type_value.lower())
            if dup_key in existing:
                status = "skipped"
                warnings.append("Already in this project")
            elif dup_key in seen:
                status = "skipped"
                warnings.append(f"Same name and type as row {seen[dup_key]} in this file")
            else:
                seen[dup_key] = sheet_row

        results.append({
            "row": sheet_row,
            "name": name,
            "type": type_value,
            "start_dt": start_dt,
            "end_dt": end_dt,
            "required_stages": chosen_stages,
            "approval_types": approvals,
            "status": status,
            "errors": errors,
            "warnings": warnings,
        })

    counts = {
        "total": len(results),
        "ok": sum(1 for r in results if r["status"] == "ok"),
        "skipped": sum(1 for r in results if r["status"] == "skipped"),
        "errors": sum(1 for r in results if r["status"] == "error"),
    }
    return {"rows": results, "counts": counts, "ignored_columns": ignored, "columns_found": sorted(columns.keys())}